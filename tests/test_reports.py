import io
from datetime import date, datetime, timedelta
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.template.context import BaseContext

# Python 3.14 compatibility patch for Django Context copy
if not getattr(BaseContext, '_copy_patched', False):
    def _safe_context_copy(self):
        duplicate = self.__class__.__new__(self.__class__)
        duplicate.__dict__.update(self.__dict__)
        if hasattr(self, 'dicts'):
            duplicate.dicts = self.dicts[:]
        return duplicate
    BaseContext.__copy__ = _safe_context_copy
    BaseContext._copy_patched = True

from accounts.models import User
from companies.models import Company, Department, Position
from employees.models import Employee
from attendance.models import Attendance
from payroll.models import Payroll, SalaryStructure


class ReportsAndExportsIntegrationTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name="Global Acme Corp", address="123 Corporate Blvd", phone="+1-555-0199")
        self.dept_eng = Department.objects.create(company=self.company, name="Engineering")
        self.dept_hr = Department.objects.create(company=self.company, name="Human Resources")
        self.pos_dev = Position.objects.create(company=self.company, department=self.dept_eng, name="Senior Engineer")
        self.pos_hr = Position.objects.create(company=self.company, department=self.dept_hr, name="HR Specialist")

        self.admin_user = User.objects.create_superuser(username="superadmin", password="password123", role=User.Role.SUPER_ADMIN)
        self.emp_user1 = User.objects.create_user(username="alice", password="password123", role=User.Role.EMPLOYEE)
        self.emp_user2 = User.objects.create_user(username="bob", password="password123", role=User.Role.EMPLOYEE)

        self.emp1 = Employee.objects.create(
            user=self.emp_user1,
            company=self.company,
            department=self.dept_eng,
            position=self.pos_dev,
            employee_code="EMP-101",
            first_name="Alice",
            last_name="Smith",
            email="alice@acme.com",
            join_date=date(2023, 1, 15),
            basic_salary=4500.00
        )
        self.emp2 = Employee.objects.create(
            user=self.emp_user2,
            company=self.company,
            department=self.dept_hr,
            position=self.pos_hr,
            employee_code="EMP-102",
            first_name="Bob",
            last_name="Jones",
            email="bob@acme.com",
            join_date=date(2023, 3, 1),
            basic_salary=3800.00
        )

        today = date.today()
        # Today's attendances
        self.att1 = Attendance.objects.create(
            employee=self.emp1,
            date=today,
            check_in=datetime.combine(today, datetime.min.time()) + timedelta(hours=8, minutes=0),
            check_out=datetime.combine(today, datetime.min.time()) + timedelta(hours=17, minutes=0),
            working_hours=timedelta(hours=9),
            status=Attendance.Status.PRESENT
        )
        self.att2 = Attendance.objects.create(
            employee=self.emp2,
            date=today,
            check_in=datetime.combine(today, datetime.min.time()) + timedelta(hours=8, minutes=45),
            check_out=datetime.combine(today, datetime.min.time()) + timedelta(hours=17, minutes=30),
            working_hours=timedelta(hours=8, minutes=45),
            status=Attendance.Status.LATE
        )

        # Payroll records
        period = date(today.year, today.month, 1)
        self.payroll1 = Payroll.objects.create(
            employee=self.emp1,
            payroll_period=period,
            basic_salary=4500.00,
            allowance=300.00,
            overtime=150.00,
            bonus=200.00,
            gross_salary=5150.00,
            tax=450.00,
            nssf=200.00,
            other_deduction=50.00,
            total_deduction=700.00,
            net_salary=4450.00,
            status=Payroll.Status.PAID
        )
        self.payroll2 = Payroll.objects.create(
            employee=self.emp2,
            payroll_period=period,
            basic_salary=3800.00,
            allowance=200.00,
            overtime=0.00,
            bonus=0.00,
            gross_salary=4000.00,
            tax=300.00,
            nssf=150.00,
            other_deduction=0.00,
            total_deduction=450.00,
            net_salary=3550.00,
            status=Payroll.Status.APPROVED
        )

        self.client = Client()

    # -------------------------------------------------------------------------
    # DAILY REPORT TESTS
    # -------------------------------------------------------------------------
    def test_daily_report_view(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:daily'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Daily HR &amp; Attendance Report")
        self.assertContains(resp, "EMP-101")
        self.assertContains(resp, "Alice Smith")

    def test_daily_report_excel_export(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:daily_excel'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertTrue(resp.content.startswith(b'PK\x03\x04'))
        
        # Verify valid openpyxl workbook
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertTrue(len(wb.sheetnames) >= 1)
        ws = wb.active
        self.assertIn("Daily HR & Attendance Report", str(ws.cell(row=2, column=1).value))

    def test_daily_report_pdf_export(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:daily_pdf'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF-'))
        self.assertTrue(len(resp.content) > 1000)

    # -------------------------------------------------------------------------
    # SUMMARY REPORT TESTS
    # -------------------------------------------------------------------------
    def test_summary_report_view(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:summary'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Executive Summary Report")
        self.assertContains(resp, "Engineering")
        self.assertContains(resp, "Human Resources")

    def test_summary_report_excel_export(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:summary_excel'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertTrue(resp.content.startswith(b'PK\x03\x04'))

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Executive Summary", wb.sheetnames)
        self.assertIn("Payroll Register", wb.sheetnames)

    def test_summary_report_pdf_export(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:summary_pdf'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF-'))
        self.assertTrue(len(resp.content) > 1000)

    # -------------------------------------------------------------------------
    # DETAIL REPORT TESTS
    # -------------------------------------------------------------------------
    def test_detail_report_view_tabs(self):
        self.client.force_login(self.admin_user)
        for tab in ['employees', 'payroll', 'attendance']:
            resp = self.client.get(reverse('reports:detail') + f'?tab={tab}')
            self.assertEqual(resp.status_code, 200)
            self.assertContains(resp, "Detailed Employee &amp; Payroll Master Ledger")

    def test_detail_report_excel_export(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('reports:detail_excel'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(set(wb.sheetnames), {"Employee Master Roster", "Payroll Register", "Attendance Ledger"})

    def test_detail_report_pdf_export_all_tabs(self):
        self.client.force_login(self.admin_user)
        for tab in ['employees', 'payroll', 'attendance']:
            resp = self.client.get(reverse('reports:detail_pdf') + f'?tab={tab}')
            self.assertEqual(resp.status_code, 200)
            self.assertEqual(resp['Content-Type'], 'application/pdf')
            self.assertTrue(resp.content.startswith(b'%PDF-'))

    # -------------------------------------------------------------------------
    # OFFICIAL PAYSLIP PDF TESTS
    # -------------------------------------------------------------------------
    def test_payslip_pdf_download_admin(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse('payroll:payslip_pdf', kwargs={'pk': self.payroll1.pk}))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF-'))
        self.assertIn('attachment;', resp['Content-Disposition'])

    def test_payslip_pdf_download_own_employee(self):
        self.client.force_login(self.emp_user1)
        resp = self.client.get(reverse('payroll:payslip_pdf', kwargs={'pk': self.payroll1.pk}))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')

    def test_payslip_pdf_download_unauthorized_employee(self):
        self.client.force_login(self.emp_user2)  # Bob trying to download Alice's payslip
        resp = self.client.get(reverse('payroll:payslip_pdf', kwargs={'pk': self.payroll1.pk}))
        self.assertEqual(resp.status_code, 404)

    # -------------------------------------------------------------------------
    # ATTENDANCE & PAYROLL LIST EXPORT TESTS
    # -------------------------------------------------------------------------
    def test_attendance_report_exports(self):
        self.client.force_login(self.admin_user)
        excel_resp = self.client.get(reverse('attendance:export_excel'))
        self.assertEqual(excel_resp.status_code, 200)
        self.assertEqual(excel_resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        pdf_resp = self.client.get(reverse('attendance:export_pdf'))
        self.assertEqual(pdf_resp.status_code, 200)
        self.assertEqual(pdf_resp['Content-Type'], 'application/pdf')

    def test_payroll_list_exports(self):
        self.client.force_login(self.admin_user)
        excel_resp = self.client.get(reverse('payroll:export_excel'))
        self.assertEqual(excel_resp.status_code, 200)
        self.assertEqual(excel_resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        pdf_resp = self.client.get(reverse('payroll:export_pdf'))
        self.assertEqual(pdf_resp.status_code, 200)
        self.assertEqual(pdf_resp['Content-Type'], 'application/pdf')

    # -------------------------------------------------------------------------
    # ACCESS CONTROL TESTS
    # -------------------------------------------------------------------------
    def test_unauthenticated_user_redirected(self):
        resp = self.client.get(reverse('reports:daily'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse('accounts:login'), resp['Location'])
