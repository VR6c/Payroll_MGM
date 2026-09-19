import os
import io
import datetime
import html
from decimal import Decimal

from django.conf import settings

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch, mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, KeepTogether, HRFlowable, Image as RLImage
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY


# ==============================================================================
# A4 PAGE DIMENSIONS (in points; A4 = 595.27 × 841.89)
# ==============================================================================
A4_WIDTH = A4[0]
A4_HEIGHT = A4[1]
A4_LANDSCAPE = landscape(A4)

MARGIN_NARROW = 30
MARGIN_STANDARD = 40

A4_PORTRAIT_USABLE = A4_WIDTH - (MARGIN_STANDARD * 2)    # ≈ 515.27
A4_LANDSCAPE_USABLE = A4_HEIGHT - (MARGIN_NARROW * 2)    # ≈ 781.89


# ==============================================================================
# PROFESSIONAL COLOR SYSTEM
# ==============================================================================

class BrandColors:
    """Centralized corporate color palette for consistent branding."""
    NAVY = colors.HexColor('#0F172A')
    NAVY_LIGHT = colors.HexColor('#1E293B')
    BLUE = colors.HexColor('#2563EB')
    BLUE_LIGHT = colors.HexColor('#3B82F6')
    BLUE_PALE = colors.HexColor('#EFF6FF')

    SUCCESS = colors.HexColor('#059669')
    SUCCESS_PALE = colors.HexColor('#ECFDF5')
    WARNING = colors.HexColor('#D97706')
    WARNING_PALE = colors.HexColor('#FFFBEB')
    DANGER = colors.HexColor('#DC2626')
    DANGER_PALE = colors.HexColor('#FEF2F2')
    INFO = colors.HexColor('#0284C7')
    INFO_PALE = colors.HexColor('#F0F9FF')

    SLATE_900 = colors.HexColor('#0F172A')
    SLATE_700 = colors.HexColor('#334155')
    SLATE_600 = colors.HexColor('#475569')
    SLATE_500 = colors.HexColor('#64748B')
    SLATE_400 = colors.HexColor('#94A3B8')
    SLATE_300 = colors.HexColor('#CBD5E1')
    SLATE_200 = colors.HexColor('#E2E8F0')
    SLATE_100 = colors.HexColor('#F1F5F9')
    SLATE_50 = colors.HexColor('#F8FAFC')
    WHITE = colors.white


# ==============================================================================
# NUMBERED CANVAS — Professional Header/Footer (A4-aware, light theme)
# ==============================================================================

class ProfessionalCanvas(canvas.Canvas):
    """
    Two-pass canvas producing executive-grade document control.
    Light theme: white header, subtle accents, blue accent footer line.
    """
    def __init__(self, *args, **kwargs):
        self._document_title = kwargs.pop('document_title', 'Report')
        self._document_ref = kwargs.pop('document_ref', '')
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self._draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def _draw_page_decorations(self, page_count):
        self.saveState()
        page_w, page_h = self._pagesize

        if page_w > page_h:
            margin = MARGIN_NARROW
        else:
            margin = MARGIN_STANDARD

        # ── Running Header (skip on first page) ────────────────────────
        if self._pageNumber > 1:
            self.setFillColor(BrandColors.WHITE)
            self.rect(0, page_h - 34, page_w, 34, stroke=0, fill=1)

            self.setStrokeColor(BrandColors.SLATE_200)
            self.setLineWidth(0.5)
            self.line(margin, page_h - 34, page_w - margin, page_h - 34)
            self.line(margin, page_h - 4, page_w - margin, page_h - 4)

            self.setFont("Helvetica-Bold", 8)
            self.setFillColor(BrandColors.SLATE_700)
            self.drawString(margin, page_h - 22, self._document_title.upper())

            self.setFont("Helvetica", 7.5)
            self.setFillColor(BrandColors.SLATE_500)
            if self._document_ref:
                self.drawRightString(page_w - margin, page_h - 22, self._document_ref)

        # ── Running Footer ──────────────────────────────────────────────
        footer_y = 26

        self.setStrokeColor(BrandColors.BLUE)
        self.setLineWidth(1.5)
        self.line(margin, footer_y + 14, margin + 60, footer_y + 14)
        self.setStrokeColor(BrandColors.SLATE_200)
        self.setLineWidth(0.5)
        self.line(margin + 60, footer_y + 14, page_w - margin, footer_y + 14)

        self.setFont("Helvetica", 7.5)
        self.setFillColor(BrandColors.SLATE_500)
        self.drawString(margin, footer_y, "CONFIDENTIAL — Payroll Management System")

        self.setFont("Helvetica", 7.5)
        self.setFillColor(BrandColors.SLATE_400)
        now_str = datetime.datetime.now().strftime("%d %b %Y, %H:%M")
        self.drawCentredString(page_w / 2.0, footer_y, f"Generated: {now_str}")

        self.setFont("Helvetica-Bold", 7.5)
        self.setFillColor(BrandColors.SLATE_600)
        self.drawRightString(page_w - margin, footer_y, f"Page {self._pageNumber} of {page_count}")

        self.restoreState()


NumberedCanvas = ProfessionalCanvas


# ==============================================================================
# PROFESSIONAL TYPOGRAPHY SYSTEM
# ==============================================================================

def get_pdf_styles():
    """Professional paragraph style registry."""
    styles = getSampleStyleSheet()

    custom_styles = {
        'DocCompany': ParagraphStyle(
            'DocCompany', fontName='Helvetica-Bold', fontSize=9,
            leading=11, textColor=BrandColors.BLUE, spaceAfter=2,
        ),
        'DocTitle': ParagraphStyle(
            'DocTitle', fontName='Helvetica-Bold', fontSize=19,
            leading=23, textColor=BrandColors.SLATE_900, spaceAfter=2,
        ),
        'DocSubtitle': ParagraphStyle(
            'DocSubtitle', fontName='Helvetica', fontSize=9,
            leading=12, textColor=BrandColors.SLATE_500, spaceAfter=0,
        ),
        'DocMeta': ParagraphStyle(
            'DocMeta', fontName='Helvetica', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_400,
        ),
        'SectionHeader': ParagraphStyle(
            'SectionHeader', fontName='Helvetica-Bold', fontSize=11,
            leading=14, textColor=BrandColors.SLATE_900,
            spaceBefore=8, spaceAfter=6,
        ),
        'SectionHeaderLight': ParagraphStyle(
            'SectionHeaderLight', fontName='Helvetica-Bold', fontSize=9,
            leading=12, textColor=BrandColors.SLATE_600,
            spaceBefore=6, spaceAfter=4,
        ),
        'TableHeader': ParagraphStyle(
            'TableHeader', fontName='Helvetica-Bold', fontSize=7.5,
            leading=10, textColor=BrandColors.WHITE, alignment=TA_LEFT,
        ),
        'TableHeaderCenter': ParagraphStyle(
            'TableHeaderCenter', fontName='Helvetica-Bold', fontSize=7.5,
            leading=10, textColor=BrandColors.WHITE, alignment=TA_CENTER,
        ),
        'TableHeaderRight': ParagraphStyle(
            'TableHeaderRight', fontName='Helvetica-Bold', fontSize=7.5,
            leading=10, textColor=BrandColors.WHITE, alignment=TA_RIGHT,
        ),
        'TableCell': ParagraphStyle(
            'TableCell', fontName='Helvetica', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_700, alignment=TA_LEFT,
        ),
        'TableCellBold': ParagraphStyle(
            'TableCellBold', fontName='Helvetica-Bold', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_900, alignment=TA_LEFT,
        ),
        'TableCellCenter': ParagraphStyle(
            'TableCellCenter', fontName='Helvetica', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_700, alignment=TA_CENTER,
        ),
        'TableCellRight': ParagraphStyle(
            'TableCellRight', fontName='Helvetica', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_700, alignment=TA_RIGHT,
        ),
        'TableCellRightBold': ParagraphStyle(
            'TableCellRightBold', fontName='Helvetica-Bold', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_900, alignment=TA_RIGHT,
        ),
        'TableCellMono': ParagraphStyle(
            'TableCellMono', fontName='Courier', fontSize=7.5,
            leading=10, textColor=BrandColors.SLATE_700, alignment=TA_LEFT,
        ),
        'BadgeSuccess': ParagraphStyle(
            'BadgeSuccess', fontName='Helvetica-Bold', fontSize=7,
            leading=9, textColor=BrandColors.SUCCESS, alignment=TA_CENTER,
        ),
        'BadgeWarning': ParagraphStyle(
            'BadgeWarning', fontName='Helvetica-Bold', fontSize=7,
            leading=9, textColor=BrandColors.WARNING, alignment=TA_CENTER,
        ),
        'BadgeDanger': ParagraphStyle(
            'BadgeDanger', fontName='Helvetica-Bold', fontSize=7,
            leading=9, textColor=BrandColors.DANGER, alignment=TA_CENTER,
        ),
        'BadgeInfo': ParagraphStyle(
            'BadgeInfo', fontName='Helvetica-Bold', fontSize=7,
            leading=9, textColor=BrandColors.INFO, alignment=TA_CENTER,
        ),
        'KPILabel': ParagraphStyle(
            'KPILabel', fontName='Helvetica-Bold', fontSize=6.5,
            leading=8, textColor=BrandColors.SLATE_500, alignment=TA_CENTER,
        ),
        'KPIValue': ParagraphStyle(
            'KPIValue', fontName='Helvetica-Bold', fontSize=13,
            leading=15, textColor=BrandColors.SLATE_900, alignment=TA_CENTER,
        ),
        'KPISub': ParagraphStyle(
            'KPISub', fontName='Helvetica', fontSize=6,
            leading=8, textColor=BrandColors.SLATE_400, alignment=TA_CENTER,
        ),
        'Notice': ParagraphStyle(
            'Notice', fontName='Helvetica-Oblique', fontSize=6.5,
            leading=9, textColor=BrandColors.SLATE_400, alignment=TA_CENTER,
        ),
    }

    for name, style in custom_styles.items():
        if name not in styles:
            styles.add(style)

    return styles


def section_divider(width=100, color=None, thickness=2):
    """Creates a branded section divider."""
    return HRFlowable(
        width=width, thickness=thickness,
        color=color or BrandColors.BLUE,
        spaceBefore=2, spaceAfter=8,
        hAlign='LEFT'
    )


def build_signature_block(page_width=500, signatories=None):
    """Professional executive authorization block."""
    styles = get_pdf_styles()
    signatories = signatories or [
        ("Prepared By", "Payroll Officer"),
        ("Verified By", "HR Manager"),
        ("Approved By", "Finance Director"),
    ]

    header_cells = []
    for title, role in signatories:
        header_cells.append([
            Paragraph(f"<b>{title}</b>", styles['TableCellBold']),
            Spacer(1, 2),
            Paragraph(f"<font size=6.5 color='#64748B'>{role}</font>", styles['TableCell']),
        ])

    sig_data = [header_cells]

    sig_data.append([
        Paragraph("<br/><br/>_________________________", styles['TableCell'])
        for _ in signatories
    ])

    sig_data.append([
        Paragraph("<font size=6 color='#94A3B8'>Signature</font>", styles['TableCellCenter'])
        for _ in signatories
    ])

    col_w = page_width / float(len(signatories))
    t = Table(sig_data, colWidths=[col_w] * len(signatories))
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, -1), (-1, -1), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('LINEBELOW', (0, 0), (-1, 0), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
    ]))
    return KeepTogether([
        Spacer(1, 16),
        HRFlowable(width="100%", thickness=0.5, color=BrandColors.SLATE_200),
        Spacer(1, 12),
        t
    ])


def build_kpi_card(label, value, sub=None, accent_color=None):
    """Builds a professional KPI card as a flowable (light theme)."""
    styles = get_pdf_styles()
    accent = accent_color or BrandColors.BLUE

    content = [
        Paragraph(label.upper(), styles['KPILabel']),
        Spacer(1, 3),
        Paragraph(f"<font color='{accent.hexval()}'>{value}</font>", styles['KPIValue']),
    ]
    if sub:
        content.append(Spacer(1, 2))
        content.append(Paragraph(sub, styles['KPISub']))

    return content


def build_status_badge(status_text, status_type='info'):
    """Renders a colored status badge."""
    styles = get_pdf_styles()
    style_map = {
        'success': 'BadgeSuccess',
        'warning': 'BadgeWarning',
        'danger': 'BadgeDanger',
        'info': 'BadgeInfo',
    }
    return Paragraph(status_text.upper(), styles[style_map.get(status_type, 'BadgeInfo')])


# ==============================================================================
# OPENPYXL STYLES & BUILDERS
# ==============================================================================

class ExcelTheme:
    NAVY_HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    TOTALS_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    SUCCESS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    DANGER_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="0F172A")
    FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="64748B")
    FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
    FONT_REGULAR = Font(name="Calibri", size=10, color="1E293B")
    FONT_TOTAL = Font(name="Calibri", size=11, bold=True, color="0F172A")

    BORDER_THIN = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    BORDER_TOTAL = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='0F172A'),
        bottom=Side(style='double', color='0F172A')
    )

    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
    ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

    FMT_CURRENCY = '$#,##0.00'
    FMT_INTEGER = '#,##0'
    FMT_PERCENT = '0.0%'
    FMT_DATE = 'yyyy-mm-dd'
    FMT_TIME = 'hh:mm'


def auto_fit_columns(ws, min_width=12, padding=4):
    """Dynamically sizes Excel columns so text is never truncated."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                lines = val.split('\n')
                line_len = max(len(l) for l in lines) if lines else 0
                max_len = max(max_len, line_len)
            else:
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + padding, min_width)


def add_excel_header_block(ws, company_name, report_title, subtitle,
                            generated_by="System", max_col=7):
    """Renders a corporate banner and metadata card on rows 1–4."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    c1 = ws.cell(row=1, column=1, value=(company_name or "Payroll Management System").upper())
    c1.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    c1.alignment = ExcelTheme.ALIGN_LEFT
    ws.row_dimensions[1].height = 18

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    c2 = ws.cell(row=2, column=1, value=report_title)
    c2.font = ExcelTheme.FONT_TITLE
    c2.alignment = ExcelTheme.ALIGN_LEFT
    ws.row_dimensions[2].height = 24

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    meta_text = f"{subtitle} | Generated on: {now_str} | Generated by: {generated_by}"
    c3 = ws.cell(row=3, column=1, value=meta_text)
    c3.font = ExcelTheme.FONT_SUBTITLE
    c3.alignment = ExcelTheme.ALIGN_LEFT
    ws.row_dimensions[3].height = 18

    ws.row_dimensions[4].height = 8


# ==============================================================================
# 1. DAILY ATTENDANCE REPORT (EXCEL & PDF — A4 Landscape)
# ==============================================================================

def generate_daily_attendance_excel(report_date, attendances, stats,
                                     company_name="Payroll MGM", generated_by="Admin"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Daily_{report_date.strftime('%Y%m%d')}"
    ws.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws,
        company_name=company_name,
        report_title="Daily HR & Attendance Report",
        subtitle=f"Date: {report_date.strftime('%B %d, %Y (%A)')}",
        generated_by=generated_by,
        max_col=7
    )

    kpis = [
        ("Total Active Employees", stats.get('total_employees', 0), ExcelTheme.FMT_INTEGER),
        ("Present", stats.get('present_count', 0), ExcelTheme.FMT_INTEGER),
        ("Late Check-ins", stats.get('late_count', 0), ExcelTheme.FMT_INTEGER),
        ("Absent", stats.get('absent_count', 0), ExcelTheme.FMT_INTEGER),
        ("Early Leave", stats.get('early_leave_count', 0), ExcelTheme.FMT_INTEGER),
        ("Overtime", stats.get('overtime_count', 0), ExcelTheme.FMT_INTEGER),
        ("Attendance Rate", stats.get('attendance_rate', 0) / 100.0, ExcelTheme.FMT_PERCENT)
    ]

    for c_idx, (kpi_label, val, fmt) in enumerate(kpis, start=1):
        cell_lbl = ws.cell(row=5, column=c_idx, value=kpi_label)
        cell_lbl.font = Font(name="Calibri", size=9, bold=True, color="475569")
        cell_lbl.fill = ExcelTheme.SUBHEADER_FILL
        cell_lbl.alignment = ExcelTheme.ALIGN_CENTER
        cell_lbl.border = ExcelTheme.BORDER_THIN

        cell_val = ws.cell(row=6, column=c_idx, value=val)
        cell_val.font = Font(name="Calibri", size=13, bold=True, color="0F172A")
        cell_val.fill = ExcelTheme.SUBHEADER_FILL
        cell_val.alignment = ExcelTheme.ALIGN_CENTER
        cell_val.border = ExcelTheme.BORDER_THIN
        cell_val.number_format = fmt

    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 24
    ws.row_dimensions[7].height = 10

    headers = ["#", "Employee Code", "Employee Name", "Department", "Position",
               "Check In", "Check Out", "Working Hours", "Status"]
    ws.row_dimensions[8].height = 26

    for c_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=8, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    curr_row = 9
    for idx, a in enumerate(attendances, start=1):
        emp = a.employee
        cin = a.check_in.strftime("%H:%M") if a.check_in else "-"
        cout = a.check_out.strftime("%H:%M") if a.check_out else "-"
        hours_str = str(a.working_hours).split('.')[0] if a.working_hours else "-"

        row_vals = [
            (idx, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (emp.position.name if emp.position else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (cin, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (cout, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (hours_str, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (a.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]

        zebra = ExcelTheme.ZEBRA_FILL if curr_row % 2 == 0 else PatternFill(fill_type=None)
        ws.row_dimensions[curr_row].height = 20

        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=curr_row, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            if c_idx == 9:
                if a.status == 'present':
                    cell.fill = ExcelTheme.SUCCESS_FILL
                elif a.status == 'late':
                    cell.fill = ExcelTheme.WARNING_FILL
                elif a.status == 'absent':
                    cell.fill = ExcelTheme.DANGER_FILL
                else:
                    cell.fill = zebra
            else:
                cell.fill = zebra
            if fmt:
                cell.number_format = fmt

        curr_row += 1

    if not attendances:
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=len(headers))
        empty_c = ws.cell(row=curr_row, column=1, value="No attendance records found for this date.")
        empty_c.font = ExcelTheme.FONT_SUBTITLE
        empty_c.alignment = ExcelTheme.ALIGN_CENTER
        empty_c.border = ExcelTheme.BORDER_THIN
        ws.row_dimensions[curr_row].height = 25

    ws.freeze_panes = 'A9'
    auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_daily_attendance_pdf(report_date, attendances, stats,
                                   company_name="Payroll MGM", generated_by="Admin"):
    """Executive-grade daily attendance report — A4 Landscape."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4_LANDSCAPE,
        leftMargin=MARGIN_NARROW, rightMargin=MARGIN_NARROW,
        topMargin=44, bottomMargin=44,
        title=f"Daily Attendance Report - {report_date.strftime('%Y-%m-%d')}",
        author=company_name,
    )

    styles = get_pdf_styles()
    story = []
    PAGE_W = A4_LANDSCAPE_USABLE

    date_str = report_date.strftime("%B %d, %Y")
    day_str = report_date.strftime("%A")

    header_left = [
        Paragraph(html.escape(company_name.upper()), styles['DocCompany']),
        Spacer(1, 3),
        Paragraph("Daily Attendance Report", styles['DocTitle']),
        Spacer(1, 2),
        Paragraph(f"<b>{date_str}</b> &nbsp;•&nbsp; {day_str}", styles['DocSubtitle']),
    ]

    header_right = [
        Paragraph("<font size=7 color='#64748B'>DOCUMENT REF</font>", styles['DocMeta']),
        Spacer(1, 2),
        Paragraph(f"<font size=9 color='#0F172A'><b>ATT-{report_date.strftime('%Y%m%d')}</b></font>", styles['DocMeta']),
        Spacer(1, 8),
        Paragraph("<font size=7 color='#64748B'>GENERATED BY</font>", styles['DocMeta']),
        Spacer(1, 2),
        Paragraph(f"<font size=9 color='#0F172A'><b>{html.escape(generated_by)}</b></font>", styles['DocMeta']),
    ]

    t_header = Table(
        [[header_left, header_right]],
        colWidths=[PAGE_W * 0.65, PAGE_W * 0.35]
    )
    t_header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=2, color=BrandColors.BLUE))
    story.append(Spacer(1, 12))

    kpi_config = [
        ("Active Workforce", str(stats.get('total_employees', 0)), "total employees", BrandColors.BLUE),
        ("Present", str(stats.get('present_count', 0)), f"{stats.get('attendance_rate', 0):.1f}% rate", BrandColors.SUCCESS),
        ("Late Check-ins", str(stats.get('late_count', 0)), "arrivals", BrandColors.WARNING),
        ("Absent", str(stats.get('absent_count', 0)), "no show", BrandColors.DANGER),
        ("Early Leave", str(stats.get('early_leave_count', 0)), "departures", BrandColors.INFO),
        ("Overtime", str(stats.get('overtime_count', 0)), "extra hours", colors.HexColor('#7C3AED')),
    ]

    kpi_cells = []
    for label, value, sub, color in kpi_config:
        kpi_cells.append(build_kpi_card(label, value, sub, color))

    kpi_table = Table([kpi_cells], colWidths=[PAGE_W / 6.0] * 6)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BrandColors.WHITE),
        ('BOX', (0, 0), (-1, -1), 0.75, BrandColors.SLATE_200),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEABOVE', (0, 0), (-1, 0), 2, BrandColors.BLUE),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 14))

    story.append(Paragraph("Employee Attendance Log", styles['SectionHeader']))
    story.append(section_divider(width=80))

    table_data = [[
        Paragraph("#", styles['TableHeaderCenter']),
        Paragraph("EMPLOYEE CODE", styles['TableHeader']),
        Paragraph("EMPLOYEE NAME", styles['TableHeader']),
        Paragraph("DEPARTMENT", styles['TableHeader']),
        Paragraph("POSITION", styles['TableHeader']),
        Paragraph("CHECK IN", styles['TableHeaderCenter']),
        Paragraph("CHECK OUT", styles['TableHeaderCenter']),
        Paragraph("HOURS", styles['TableHeaderCenter']),
        Paragraph("STATUS", styles['TableHeaderCenter']),
    ]]

    for idx, a in enumerate(attendances, start=1):
        emp = a.employee
        cin = a.check_in.strftime("%H:%M") if a.check_in else "—"
        cout = a.check_out.strftime("%H:%M") if a.check_out else "—"
        hours_str = str(a.working_hours).split('.')[0] if a.working_hours else "—"

        status_type = 'success' if a.status == 'present' else (
            'warning' if a.status == 'late' else (
                'danger' if a.status == 'absent' else 'info'
            )
        )

        table_data.append([
            Paragraph(str(idx), styles['TableCellCenter']),
            Paragraph(html.escape(emp.employee_code), styles['TableCellMono']),
            Paragraph(f"<b>{html.escape(emp.first_name)} {html.escape(emp.last_name)}</b>", styles['TableCell']),
            Paragraph(html.escape(emp.department.name if emp.department else "—"), styles['TableCell']),
            Paragraph(html.escape(emp.position.name if emp.position else "—"), styles['TableCell']),
            Paragraph(cin, styles['TableCellCenter']),
            Paragraph(cout, styles['TableCellCenter']),
            Paragraph(hours_str, styles['TableCellCenter']),
            build_status_badge(a.get_status_display(), status_type),
        ])

    if not attendances:
        table_data.append([
            Paragraph("No attendance records found for this date.", styles['TableCellCenter']),
            "", "", "", "", "", "", "", ""
        ])

    col_widths = [30, 75, 165, 120, 120, 65, 65, 60, 81.89]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), BrandColors.NAVY),
        ('TEXTCOLOR', (0, 0), (-1, 0), BrandColors.WHITE),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [BrandColors.WHITE, BrandColors.SLATE_50]),
        ('GRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, BrandColors.BLUE),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]
    if not attendances:
        t_style.append(('SPAN', (0, 1), (-1, 1)))

    t.setStyle(TableStyle(t_style))
    story.append(t)

    story.append(build_signature_block(page_width=PAGE_W))

    doc.build(
        story,
        canvasmaker=lambda *a, **kw: ProfessionalCanvas(
            *a, **kw,
            document_title="Daily Attendance Report",
            document_ref=f"ATT-{report_date.strftime('%Y%m%d')}"
        )
    )
    return buffer.getvalue()


# ==============================================================================
# 2. EXECUTIVE SUMMARY REPORT (EXCEL & PDF — A4 Portrait)
# ==============================================================================

def generate_summary_report_excel(period_str, payrolls, dept_stats, summary_stats,
                                   company_name="Payroll MGM", generated_by="Admin"):
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws1,
        company_name=company_name,
        report_title="Executive Payroll & HR Summary Report",
        subtitle=f"Period: {period_str or 'All Periods'}",
        generated_by=generated_by,
        max_col=6
    )

    tiles = [
        ("Total Workforce", summary_stats.get('total_employees', 0), ExcelTheme.FMT_INTEGER),
        ("Active Employees", summary_stats.get('active_employees', 0), ExcelTheme.FMT_INTEGER),
        ("Total Gross Payroll", summary_stats.get('total_gross', 0), ExcelTheme.FMT_CURRENCY),
        ("Total Deductions", summary_stats.get('total_deductions', 0), ExcelTheme.FMT_CURRENCY),
        ("Net Payroll Paid", summary_stats.get('total_net', 0), ExcelTheme.FMT_CURRENCY),
        ("Average Net Salary", summary_stats.get('avg_salary', 0), ExcelTheme.FMT_CURRENCY),
    ]

    for c_idx, (lbl, val, fmt) in enumerate(tiles, start=1):
        c_lbl = ws1.cell(row=5, column=c_idx, value=lbl)
        c_lbl.font = Font(name="Calibri", size=9, bold=True, color="475569")
        c_lbl.fill = ExcelTheme.SUBHEADER_FILL
        c_lbl.alignment = ExcelTheme.ALIGN_CENTER
        c_lbl.border = ExcelTheme.BORDER_THIN

        c_val = ws1.cell(row=6, column=c_idx, value=val)
        c_val.font = Font(name="Calibri", size=13, bold=True, color="0F172A")
        c_val.fill = ExcelTheme.SUBHEADER_FILL
        c_val.alignment = ExcelTheme.ALIGN_CENTER
        c_val.border = ExcelTheme.BORDER_THIN
        c_val.number_format = fmt

    ws1.row_dimensions[5].height = 18
    ws1.row_dimensions[6].height = 24
    ws1.row_dimensions[7].height = 12

    ws1.cell(row=8, column=1, value="Department Payroll & Workforce Distribution").font = ExcelTheme.FONT_TITLE
    ws1.row_dimensions[8].height = 22

    dept_headers = ["Department", "Headcount", "Total Gross Salary",
                    "Total Deductions", "Total Net Salary", "% Share"]
    ws1.row_dimensions[9].height = 24
    for c_idx, h in enumerate(dept_headers, start=1):
        c = ws1.cell(row=9, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    row_idx = 10
    total_net_sum = float(summary_stats.get('total_net', 0)) or 1.0

    for d in dept_stats:
        net_val = float(d.get('total_net', 0))
        share = net_val / total_net_sum if total_net_sum > 0 else 0

        ws1.row_dimensions[row_idx].height = 20
        vals = [
            (d.get('name', 'Unassigned'), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (d.get('count', 0), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_INTEGER),
            (float(d.get('total_gross', 0)), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(d.get('total_deductions', 0)), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (net_val, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (share, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_PERCENT),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(vals, start=1):
            cell = ws1.cell(row=row_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt
        row_idx += 1

    ws1.row_dimensions[row_idx].height = 22
    tot_vals = [
        ("Total / Grand Summary", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_TOTAL, None),
        (f"=SUM(B10:B{row_idx-1})" if row_idx > 10 else 0, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_INTEGER),
        (f"=SUM(C10:C{row_idx-1})" if row_idx > 10 else 0, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
        (f"=SUM(D10:D{row_idx-1})" if row_idx > 10 else 0, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
        (f"=SUM(E10:E{row_idx-1})" if row_idx > 10 else 0, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
        (1.0, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_PERCENT),
    ]
    for c_idx, (v, align, font, fmt) in enumerate(tot_vals, start=1):
        cell = ws1.cell(row=row_idx, column=c_idx, value=v)
        cell.font = font
        cell.alignment = align
        cell.border = ExcelTheme.BORDER_TOTAL
        cell.fill = ExcelTheme.TOTALS_FILL
        if fmt:
            cell.number_format = fmt

    auto_fit_columns(ws1)

    ws2 = wb.create_sheet(title="Payroll Register")
    ws2.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws2,
        company_name=company_name,
        report_title="Period Payroll Detailed Register",
        subtitle=f"Period: {period_str or 'All Periods'}",
        generated_by=generated_by,
        max_col=10
    )

    p_headers = ["#", "Period", "Employee Code", "Employee Name", "Department",
                 "Basic Salary", "Gross Salary", "Total Deductions", "Net Salary", "Status"]
    ws2.row_dimensions[5].height = 24

    for c_idx, h in enumerate(p_headers, start=1):
        c = ws2.cell(row=5, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    p_row = 6
    for idx, p in enumerate(payrolls, start=1):
        emp = p.employee
        ws2.row_dimensions[p_row].height = 20
        row_vals = [
            (idx, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (p.payroll_period.strftime("%b %Y"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (float(p.basic_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.gross_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.total_deduction), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.net_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (p.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if p_row % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws2.cell(row=p_row, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt
        p_row += 1

    if payrolls:
        ws2.row_dimensions[p_row].height = 22
        tots = [
            ("Grand Total", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            (f"=SUM(F6:F{p_row-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            (f"=SUM(G6:G{p_row-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            (f"=SUM(H6:H{p_row-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            (f"=SUM(I6:I{p_row-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
        ]
        for c_idx, (v, align, font, fmt) in enumerate(tots, start=1):
            cell = ws2.cell(row=p_row, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_TOTAL
            cell.fill = ExcelTheme.TOTALS_FILL
            if fmt:
                cell.number_format = fmt

    ws2.freeze_panes = 'A6'
    auto_fit_columns(ws2)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_summary_report_pdf(period_str, payrolls, dept_stats, summary_stats,
                                 company_name="Payroll MGM", generated_by="Admin"):
    """Executive-grade payroll summary report — A4 Portrait, light theme."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=MARGIN_STANDARD, rightMargin=MARGIN_STANDARD,
        topMargin=44, bottomMargin=44,
        title=f"Executive Payroll Summary - {period_str or 'All Periods'}",
        author=company_name,
    )

    styles = get_pdf_styles()
    story = []
    PAGE_W = A4_PORTRAIT_USABLE

    now_str = datetime.datetime.now().strftime("%d %b %Y, %H:%M")

    header_left = [
        Paragraph(html.escape(company_name.upper()), styles['DocCompany']),
        Spacer(1, 3),
        Paragraph("Executive Payroll Summary", styles['DocTitle']),
        Spacer(1, 2),
        Paragraph(f"Reporting Period: <b>{html.escape(period_str or 'All Periods')}</b>", styles['DocSubtitle']),
    ]

    header_right = [
        Paragraph("<font size=7 color='#64748B'>DOCUMENT REF</font>", styles['DocMeta']),
        Spacer(1, 2),
        Paragraph(f"<font size=9 color='#0F172A'><b>EXEC-{datetime.datetime.now().strftime('%Y%m%d')}</b></font>", styles['DocMeta']),
        Spacer(1, 8),
        Paragraph("<font size=7 color='#64748B'>GENERATED</font>", styles['DocMeta']),
        Spacer(1, 2),
        Paragraph(f"<font size=9 color='#0F172A'><b>{now_str}</b></font>", styles['DocMeta']),
    ]

    t_header = Table(
        [[header_left, header_right]],
        colWidths=[PAGE_W * 0.6, PAGE_W * 0.4]
    )
    t_header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=2, color=BrandColors.BLUE))
    story.append(Spacer(1, 14))

    kpi_config = [
        ("Total Workforce", f"{summary_stats.get('total_employees', 0):,}", "employees", BrandColors.BLUE),
        ("Active Employees", f"{summary_stats.get('active_employees', 0):,}", "currently active", BrandColors.SUCCESS),
        ("Gross Payroll", f"${summary_stats.get('total_gross', 0):,.2f}", "before deductions", BrandColors.SLATE_700),
    ]

    kpi_cells = []
    for label, value, sub, color in kpi_config:
        kpi_cells.append(build_kpi_card(label, value, sub, color))

    kpi_table = Table([kpi_cells], colWidths=[PAGE_W / 3.0] * 3)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BrandColors.WHITE),
        ('BOX', (0, 0), (-1, -1), 0.75, BrandColors.SLATE_200),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 12),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEABOVE', (0, 0), (-1, 0), 2, BrandColors.BLUE),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 8))

    kpi_config_2 = [
        ("Total Deductions", f"${summary_stats.get('total_deductions', 0):,.2f}", "tax & statutory", BrandColors.DANGER),
        ("Net Disbursed", f"${summary_stats.get('total_net', 0):,.2f}", "take-home pay", BrandColors.SUCCESS),
        ("Average Net Salary", f"${summary_stats.get('avg_salary', 0):,.2f}", "per employee", BrandColors.SLATE_700),
    ]

    kpi_cells_2 = []
    for label, value, sub, color in kpi_config_2:
        kpi_cells_2.append(build_kpi_card(label, value, sub, color))

    kpi_table_2 = Table([kpi_cells_2], colWidths=[PAGE_W / 3.0] * 3)
    kpi_table_2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BrandColors.WHITE),
        ('BOX', (0, 0), (-1, -1), 0.75, BrandColors.SLATE_200),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 12),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(kpi_table_2)
    story.append(Spacer(1, 18))

    story.append(Paragraph("Department Payroll Distribution", styles['SectionHeader']))
    story.append(section_divider(width=80))

    dept_table = [[
        Paragraph("DEPARTMENT", styles['TableHeader']),
        Paragraph("HEADS", styles['TableHeaderCenter']),
        Paragraph("GROSS", styles['TableHeaderRight']),
        Paragraph("DEDUCT.", styles['TableHeaderRight']),
        Paragraph("NET SALARY", styles['TableHeaderRight']),
        Paragraph("% SHARE", styles['TableHeaderRight']),
    ]]

    total_net = float(summary_stats.get('total_net', 0)) or 1.0
    for d in dept_stats:
        d_net = float(d.get('total_net', 0))
        share = (d_net / total_net * 100.0) if total_net > 0 else 0
        dept_table.append([
            Paragraph(f"<b>{html.escape(d.get('name', 'Unassigned'))}</b>", styles['TableCell']),
            Paragraph(f"{d.get('count', 0):,}", styles['TableCellCenter']),
            Paragraph(f"${d.get('total_gross', 0):,.2f}", styles['TableCellRight']),
            Paragraph(f"${d.get('total_deductions', 0):,.2f}", styles['TableCellRight']),
            Paragraph(f"<b>${d_net:,.2f}</b>", styles['TableCellRightBold']),
            Paragraph(f"{share:.1f}%", styles['TableCellRight']),
        ])

    dept_table.append([
        Paragraph("<b>TOTAL</b>", styles['TableCellBold']),
        Paragraph(f"<b>{summary_stats.get('total_employees', 0):,}</b>", styles['TableCellCenter']),
        Paragraph(f"<b>${summary_stats.get('total_gross', 0):,.2f}</b>", styles['TableCellRightBold']),
        Paragraph(f"<b>${summary_stats.get('total_deductions', 0):,.2f}</b>", styles['TableCellRightBold']),
        Paragraph(f"<b>${summary_stats.get('total_net', 0):,.2f}</b>", styles['TableCellRightBold']),
        Paragraph("<b>100.0%</b>", styles['TableCellRightBold']),
    ])

    t_dept = Table(dept_table, colWidths=[130, 45, 85, 80, 100, 75.27])
    t_dept.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), BrandColors.NAVY),
        ('BACKGROUND', (0, -1), (-1, -1), BrandColors.BLUE_PALE),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, BrandColors.BLUE),
        ('LINEBELOW', (0, -1), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [BrandColors.WHITE, BrandColors.SLATE_50]),
        ('GRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_dept)
    story.append(Spacer(1, 18))

    story.append(Paragraph("Payroll Ledger Summary", styles['SectionHeader']))
    story.append(section_divider(width=80))

    p_table = [[
        Paragraph("PERIOD", styles['TableHeader']),
        Paragraph("EMPLOYEE", styles['TableHeader']),
        Paragraph("DEPARTMENT", styles['TableHeader']),
        Paragraph("GROSS", styles['TableHeaderRight']),
        Paragraph("NET SALARY", styles['TableHeaderRight']),
        Paragraph("STATUS", styles['TableHeaderCenter']),
    ]]

    for p in payrolls[:12]:
        status_type = 'success' if p.status == 'paid' else (
            'info' if p.status == 'approved' else 'warning'
        )
        p_table.append([
            Paragraph(p.payroll_period.strftime("%b %Y"), styles['TableCell']),
            Paragraph(f"<b>{html.escape(p.employee.first_name)} {html.escape(p.employee.last_name)}</b>", styles['TableCell']),
            Paragraph(html.escape(p.employee.department.name if p.employee.department else "—"), styles['TableCell']),
            Paragraph(f"${p.gross_salary:,.2f}", styles['TableCellRight']),
            Paragraph(f"<b>${p.net_salary:,.2f}</b>", styles['TableCellRightBold']),
            build_status_badge(p.get_status_display(), status_type),
        ])

    if not payrolls:
        p_table.append([
            Paragraph("No payroll records found for this period.", styles['TableCellCenter']),
            "", "", "", "", ""
        ])

    t_p = Table(p_table, colWidths=[55, 125, 100, 80, 85, 70.27], repeatRows=1)
    p_style = [
        ('BACKGROUND', (0, 0), (-1, 0), BrandColors.NAVY),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [BrandColors.WHITE, BrandColors.SLATE_50]),
        ('GRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, BrandColors.BLUE),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]
    if not payrolls:
        p_style.append(('SPAN', (0, 1), (-1, 1)))
    t_p.setStyle(TableStyle(p_style))
    story.append(t_p)

    story.append(build_signature_block(page_width=PAGE_W))

    doc.build(
        story,
        canvasmaker=lambda *a, **kw: ProfessionalCanvas(
            *a, **kw,
            document_title="Executive Payroll Summary",
            document_ref=f"EXEC-{datetime.datetime.now().strftime('%Y%m%d')}"
        )
    )
    return buffer.getvalue()


# ==============================================================================
# 3. DETAIL REPORT (MULTI-TAB EXCEL & A4 Landscape PDF)
# ==============================================================================

def generate_detail_report_excel(employees, payrolls, attendances,
                                  company_name="Payroll MGM", generated_by="Admin"):
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Employee Master Roster"
    ws1.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws1,
        company_name=company_name,
        report_title="Employee Master Roster Ledger",
        subtitle=f"Total Records: {len(employees)}",
        generated_by=generated_by,
        max_col=10
    )

    e_headers = ["#", "Code", "First Name", "Last Name", "Gender", "Email",
                 "Phone", "Department", "Position", "Basic Salary", "Join Date", "Status"]
    ws1.row_dimensions[5].height = 24

    for c_idx, h in enumerate(e_headers, start=1):
        c = ws1.cell(row=5, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    for r_idx, emp in enumerate(employees, start=6):
        ws1.row_dimensions[r_idx].height = 20
        row_vals = [
            (r_idx - 5, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.first_name, ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.last_name, ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.get_gender_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.email, ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (emp.phone or "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (emp.position.name if emp.position else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (float(emp.basic_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (emp.join_date.strftime("%Y-%m-%d"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    ws1.freeze_panes = 'A6'
    auto_fit_columns(ws1)

    ws2 = wb.create_sheet(title="Payroll Register")
    ws2.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws2,
        company_name=company_name,
        report_title="Payroll Master Register",
        subtitle=f"Total Records: {len(payrolls)}",
        generated_by=generated_by,
        max_col=11
    )

    pr_headers = ["#", "Period", "Employee Code", "Employee Name", "Department",
                  "Basic Salary", "Allowance", "Overtime", "Bonus", "Gross Salary",
                  "Tax", "NSSF", "Other Ded.", "Total Ded.", "Net Salary", "Status"]
    ws2.row_dimensions[5].height = 24

    for c_idx, h in enumerate(pr_headers, start=1):
        c = ws2.cell(row=5, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    for r_idx, p in enumerate(payrolls, start=6):
        emp = p.employee
        ws2.row_dimensions[r_idx].height = 20
        row_vals = [
            (r_idx - 5, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (p.payroll_period.strftime("%Y-%m"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (float(p.basic_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.allowance), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.overtime), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.bonus), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.gross_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (float(p.tax), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.nssf), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.other_deduction), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.total_deduction), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.net_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            (p.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    ws2.freeze_panes = 'A6'
    auto_fit_columns(ws2)

    ws3 = wb.create_sheet(title="Attendance Ledger")
    ws3.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws3,
        company_name=company_name,
        report_title="Attendance & Time Logs",
        subtitle=f"Total Records: {len(attendances)}",
        generated_by=generated_by,
        max_col=8
    )

    att_headers = ["#", "Date", "Employee Code", "Employee Name", "Department",
                   "Check In", "Check Out", "Working Hours", "Status"]
    ws3.row_dimensions[5].height = 24

    for c_idx, h in enumerate(att_headers, start=1):
        c = ws3.cell(row=5, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    for r_idx, a in enumerate(attendances, start=6):
        emp = a.employee
        cin = a.check_in.strftime("%H:%M") if a.check_in else "-"
        cout = a.check_out.strftime("%H:%M") if a.check_out else "-"
        hours_str = str(a.working_hours).split('.')[0] if a.working_hours else "-"

        ws3.row_dimensions[r_idx].height = 20
        row_vals = [
            (r_idx - 5, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (a.date.strftime("%Y-%m-%d"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (cin, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (cout, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (hours_str, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (a.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    ws3.freeze_panes = 'A6'
    auto_fit_columns(ws3)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_detail_report_pdf(category, data_list,
                                company_name="Payroll MGM", generated_by="Admin"):
    """Landscape detailed ledger — A4 Landscape."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4_LANDSCAPE,
        leftMargin=MARGIN_NARROW, rightMargin=MARGIN_NARROW,
        topMargin=44, bottomMargin=44,
        title=f"Detailed Report - {category.title()}",
        author=company_name,
    )

    styles = get_pdf_styles()
    story = []
    PAGE_W = A4_LANDSCAPE_USABLE

    title_map = {
        'employees': "Employee Master Roster",
        'payroll': "Payroll Register Ledger",
        'attendance': "Attendance & Timesheet Ledger",
    }
    report_title = title_map.get(category, "Detailed Ledger Report")

    header_left = [
        Paragraph(html.escape(company_name.upper()), styles['DocCompany']),
        Spacer(1, 3),
        Paragraph(report_title, styles['DocTitle']),
        Spacer(1, 2),
        Paragraph(f"Comprehensive {category.title()} Records", styles['DocSubtitle']),
    ]

    header_right = [
        Paragraph("<font size=7 color='#64748B'>DOCUMENT REF</font>", styles['DocMeta']),
        Spacer(1, 2),
        Paragraph(f"<font size=9 color='#0F172A'><b>DTL-{category.upper()[:3]}-{datetime.datetime.now().strftime('%Y%m%d')}</b></font>", styles['DocMeta']),
        Spacer(1, 8),
        Paragraph("<font size=7 color='#64748B'>TOTAL RECORDS</font>", styles['DocMeta']),
        Spacer(1, 2),
        Paragraph(f"<font size=9 color='#0F172A'><b>{len(data_list):,}</b></font>", styles['DocMeta']),
    ]

    t_header = Table(
        [[header_left, header_right]],
        colWidths=[PAGE_W * 0.65, PAGE_W * 0.35]
    )
    t_header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=2, color=BrandColors.BLUE))
    story.append(Spacer(1, 12))

    if category == 'payroll':
        headers = ["#", "PERIOD", "CODE", "EMPLOYEE NAME", "DEPARTMENT", "BASIC",
                   "ALLOW.", "OT", "GROSS", "TAX", "DEDUCT.", "NET SALARY", "STATUS"]
        col_w = [26, 48, 55, 130, 88, 55, 48, 42, 60, 48, 55, 70, 57.89]
        t_data = [[
            Paragraph(h, styles['TableHeaderCenter'] if h in ["#", "PERIOD", "CODE", "STATUS"] else
                       styles['TableHeaderRight'] if h in ["BASIC", "ALLOW.", "OT", "GROSS", "TAX", "DEDUCT.", "NET SALARY"] else
                       styles['TableHeader']) for h in headers
        ]]

        for idx, p in enumerate(data_list, start=1):
            emp = p.employee
            status_type = 'success' if p.status == 'paid' else (
                'info' if p.status == 'approved' else 'warning'
            )
            t_data.append([
                Paragraph(str(idx), styles['TableCellCenter']),
                Paragraph(p.payroll_period.strftime("%b %y"), styles['TableCellCenter']),
                Paragraph(html.escape(emp.employee_code), styles['TableCellMono']),
                Paragraph(f"<b>{html.escape(emp.first_name)} {html.escape(emp.last_name)}</b>", styles['TableCell']),
                Paragraph(html.escape(emp.department.name if emp.department else "—"), styles['TableCell']),
                Paragraph(f"${p.basic_salary:,.0f}", styles['TableCellRight']),
                Paragraph(f"${p.allowance:,.0f}", styles['TableCellRight']),
                Paragraph(f"${p.overtime:,.0f}", styles['TableCellRight']),
                Paragraph(f"<b>${p.gross_salary:,.0f}</b>", styles['TableCellRightBold']),
                Paragraph(f"${p.tax:,.0f}", styles['TableCellRight']),
                Paragraph(f"-${p.total_deduction:,.0f}", styles['TableCellRight']),
                Paragraph(f"<b>${p.net_salary:,.0f}</b>", styles['TableCellRightBold']),
                build_status_badge(p.get_status_display(), status_type),
            ])
    elif category == 'attendance':
        headers = ["#", "DATE", "CODE", "EMPLOYEE NAME", "DEPARTMENT",
                   "CHECK IN", "CHECK OUT", "HOURS", "STATUS"]
        col_w = [32, 75, 68, 200, 150, 75, 75, 55, 51.89]
        t_data = [[
            Paragraph(h, styles['TableHeaderCenter'] if h in ["#", "DATE", "CODE", "CHECK IN", "CHECK OUT", "HOURS", "STATUS"] else styles['TableHeader'])
            for h in headers
        ]]

        for idx, a in enumerate(data_list, start=1):
            emp = a.employee
            cin = a.check_in.strftime("%H:%M") if a.check_in else "—"
            cout = a.check_out.strftime("%H:%M") if a.check_out else "—"
            hours_str = str(a.working_hours).split('.')[0] if a.working_hours else "—"
            status_type = 'success' if a.status == 'present' else (
                'warning' if a.status == 'late' else (
                    'danger' if a.status == 'absent' else 'info'
                )
            )
            t_data.append([
                Paragraph(str(idx), styles['TableCellCenter']),
                Paragraph(a.date.strftime("%Y-%m-%d"), styles['TableCellCenter']),
                Paragraph(html.escape(emp.employee_code), styles['TableCellMono']),
                Paragraph(f"<b>{html.escape(emp.first_name)} {html.escape(emp.last_name)}</b>", styles['TableCell']),
                Paragraph(html.escape(emp.department.name if emp.department else "—"), styles['TableCell']),
                Paragraph(cin, styles['TableCellCenter']),
                Paragraph(cout, styles['TableCellCenter']),
                Paragraph(hours_str, styles['TableCellCenter']),
                build_status_badge(a.get_status_display(), status_type),
            ])
    else:
        headers = ["#", "CODE", "EMPLOYEE NAME", "GENDER", "EMAIL", "PHONE",
                   "DEPARTMENT", "POSITION", "BASIC SALARY", "STATUS"]
        col_w = [32, 68, 160, 50, 150, 85, 90, 90, 78, 78.89]
        t_data = [[
            Paragraph(h, styles['TableHeaderRight'] if h == "BASIC SALARY" else
                       styles['TableHeaderCenter'] if h in ["#", "CODE", "GENDER", "STATUS"] else
                       styles['TableHeader']) for h in headers
        ]]

        for idx, emp in enumerate(data_list, start=1):
            status_type = 'success' if emp.status == 'active' else 'danger'
            t_data.append([
                Paragraph(str(idx), styles['TableCellCenter']),
                Paragraph(html.escape(emp.employee_code), styles['TableCellMono']),
                Paragraph(f"<b>{html.escape(emp.first_name)} {html.escape(emp.last_name)}</b>", styles['TableCellBold']),
                Paragraph(emp.get_gender_display()[:1], styles['TableCellCenter']),
                Paragraph(html.escape(emp.email), styles['TableCell']),
                Paragraph(html.escape(emp.phone or "—"), styles['TableCell']),
                Paragraph(html.escape(emp.department.name if emp.department else "—"), styles['TableCell']),
                Paragraph(html.escape(emp.position.name if emp.position else "—"), styles['TableCell']),
                Paragraph(f"<b>${emp.basic_salary:,.0f}</b>", styles['TableCellRightBold']),
                build_status_badge(emp.get_status_display(), status_type),
            ])

    if len(t_data) == 1:
        t_data.append([Paragraph("No records found.", styles['TableCellCenter'])] + [""] * (len(headers) - 1))

    t = Table(t_data, colWidths=col_w, repeatRows=1)
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), BrandColors.NAVY),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, BrandColors.BLUE),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [BrandColors.WHITE, BrandColors.SLATE_50]),
        ('GRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]
    if len(t_data) == 2 and not data_list:
        t_style.append(('SPAN', (0, 1), (-1, 1)))

    t.setStyle(TableStyle(t_style))
    story.append(t)
    story.append(build_signature_block(page_width=PAGE_W))

    doc.build(
        story,
        canvasmaker=lambda *a, **kw: ProfessionalCanvas(
            *a, **kw,
            document_title=report_title,
            document_ref=f"DTL-{category.upper()[:3]}-{datetime.datetime.now().strftime('%Y%m%d')}"
        )
    )
    return buffer.getvalue()


# ==============================================================================
# 4. OFFICIAL PAYSLIP PDF (A4 Portrait — LIGHT THEME, FIXED LAYOUT)
# ==============================================================================

def generate_payslip_pdf(payroll, company=None):
    """
    Premium, executive-grade salary payslip — A4 Portrait — LIGHT THEME.
    Clean white background, subtle borders, blue accents only.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=MARGIN_STANDARD, rightMargin=MARGIN_STANDARD,
        topMargin=30, bottomMargin=40,
        title=f"Payslip - {payroll.employee.first_name} {payroll.employee.last_name} - {payroll.payroll_period.strftime('%B %Y')}",
        author="Payroll Management System",
    )

    styles = get_pdf_styles()
    story = []
    PAGE_W = A4_PORTRAIT_USABLE

    emp = payroll.employee
    company_obj = company or getattr(emp, 'company', None)
    company_name = company_obj.name if company_obj else "Payroll MGM Inc"
    company_addr = getattr(company_obj, 'address', '') if company_obj else ""
    company_phone = getattr(company_obj, 'phone', '') if company_obj else ""
    company_email = getattr(company_obj, 'email', '') if company_obj else ""

    # ── Status palette ──────────────────────────────────────────────
    status_val = payroll.get_status_display().upper()
    status_color_map = {
        'paid': BrandColors.SUCCESS,
        'approved': BrandColors.INFO,
        'processing': BrandColors.WARNING,
        'draft': BrandColors.SLATE_500,
    }
    status_color = status_color_map.get(payroll.status, BrandColors.SLATE_500)

    # ==============================================================
    # SECTION 1: LIGHT HEADER
    # ==============================================================
    hdr_company_style = ParagraphStyle(
        'PayslipHdrCompany', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=15, leading=18,
        textColor=BrandColors.SLATE_900, alignment=TA_LEFT,
    )
    hdr_tag_style = ParagraphStyle(
        'PayslipHdrTag', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7, leading=9,
        textColor=BrandColors.BLUE, alignment=TA_RIGHT,
    )
    hdr_addr_style = ParagraphStyle(
        'PayslipHdrAddr', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7, leading=10,
        textColor=BrandColors.SLATE_500, alignment=TA_LEFT,
    )
    hdr_period_style = ParagraphStyle(
        'PayslipHdrPeriod', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=15, leading=18,
        textColor=BrandColors.SLATE_900, alignment=TA_RIGHT,
    )
    hdr_meta_style = ParagraphStyle(
        'PayslipHdrMeta', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7, leading=10,
        textColor=BrandColors.SLATE_400, alignment=TA_RIGHT,
    )

    logo_flowable = None
    logo_path = None
    if company_obj and getattr(company_obj, 'logo', None):
        try:
            if company_obj.logo and hasattr(company_obj.logo, 'path') and os.path.exists(company_obj.logo.path):
                logo_path = company_obj.logo.path
        except Exception:
            logo_path = None
    if not logo_path:
        default_logo = os.path.join(settings.BASE_DIR, 'static', 'assets', 'images', 'salary.png')
        if os.path.exists(default_logo):
            logo_path = default_logo

    if logo_path:
        try:
            logo_img = RLImage(logo_path, width=26, height=26)
            t_logo = Table([[logo_img]], colWidths=[32], rowHeights=[32])
            t_logo.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), BrandColors.WHITE),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
                ('BOX', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
            ]))
            logo_flowable = t_logo
        except Exception:
            logo_flowable = None

    addr_line = html.escape(company_addr) if company_addr else "Phnom Penh Headquarters, Cambodia"
    contact_parts = []
    if company_phone:
        contact_parts.append(f"Tel: {html.escape(company_phone)}")
    if company_email:
        contact_parts.append(html.escape(company_email))
    contact_line = "  •  ".join(contact_parts)

    brand_text = [
        Paragraph(html.escape(company_name.upper()), hdr_company_style),
        Spacer(1, 3),
        Paragraph(addr_line, hdr_addr_style),
    ]
    if contact_line:
        brand_text.append(Paragraph(contact_line, hdr_addr_style))

    left_col_w = PAGE_W * 0.60
    if logo_flowable:
        t_brand = Table([[logo_flowable, brand_text]], colWidths=[40, left_col_w - 40])
        t_brand.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        hdr_left = t_brand
    else:
        hdr_left = brand_text

    hdr_right = [
        Paragraph("OFFICIAL SALARY PAYSLIP", hdr_tag_style),
        Spacer(1, 4),
        Paragraph(payroll.payroll_period.strftime('%B %Y'), hdr_period_style),
        Spacer(1, 3),
        Paragraph(f"Ref: PS-{payroll.id:06d}  •  {payroll.created_at.strftime('%d %b %Y')}", hdr_meta_style),
    ]

    t_hdr = Table([[hdr_left, hdr_right]], colWidths=[left_col_w, PAGE_W * 0.40])
    t_hdr.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BrandColors.WHITE),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
        ('LEFTPADDING', (0, 0), (0, -1), 0),
        ('RIGHTPADDING', (-1, 0), (-1, -1), 0),
    ]))
    story.append(t_hdr)

    story.append(HRFlowable(width="100%", thickness=2, color=BrandColors.BLUE,
                             spaceBefore=0, spaceAfter=12))

    # ==============================================================
    # SECTION 2: EMPLOYEE & PAYROLL INFO (FIXED — no overflow)
    # ==============================================================
    emp_name = f"{html.escape(emp.first_name)} {html.escape(emp.last_name)}"
    emp_dept = html.escape(emp.department.name if emp.department else '—')
    emp_pos = html.escape(emp.position.name if emp.position else '—')
    emp_code = html.escape(emp.employee_code)

    # Left: employee identity
    info_left = [
        Paragraph("<font size=6.5 color='#94A3B8'><b>EMPLOYEE</b></font>", styles['TableCell']),
        Spacer(1, 3),
        Paragraph(f"<font size=13 color='#0F172A'><b>{emp_name}</b></font>", styles['TableCell']),
        Spacer(1, 4),
        Paragraph(f"<font size=8 color='#475569'>{emp_pos}</font>", styles['TableCell']),
        Spacer(1, 2),
        Paragraph(f"<font size=8 color='#64748B'>{emp_dept}</font>", styles['TableCell']),
        Spacer(1, 4),
        Paragraph(f"<font size=7 color='#94A3B8'>ID: {emp_code}</font>", styles['TableCell']),
    ]

    issue_date = payroll.created_at.strftime('%d %b %Y')
    daily_rate_str = f"${payroll.daily_salary:,.2f}/day" if payroll.daily_salary else '—'

    # Right: payroll meta (three equal thirds)
    # Right panel width = PAGE_W * 0.42, minus 12pt inner padding
    RIGHT_PANEL_W = PAGE_W * 0.42 - 12
    meta_col_w = RIGHT_PANEL_W / 3.0

    info_right_data = [
        [
            Paragraph("<font size=6.5 color='#94A3B8'><b>ISSUE DATE</b></font>", styles['TableCellCenter']),
            Paragraph("<font size=6.5 color='#94A3B8'><b>DAILY RATE</b></font>", styles['TableCellCenter']),
            Paragraph("<font size=6.5 color='#94A3B8'><b>STATUS</b></font>", styles['TableCellCenter']),
        ],
        [
            Paragraph(f"<font size=8.5 color='#0F172A'><b>{issue_date}</b></font>", styles['TableCellCenter']),
            Paragraph(f"<font size=8.5 color='#0F172A'><b>{daily_rate_str}</b></font>", styles['TableCellCenter']),
            Paragraph(f"<font size=8.5 color='{status_color.hexval()}'><b>{status_val}</b></font>", styles['TableCellCenter']),
        ],
    ]
    t_info_right = Table(info_right_data, colWidths=[meta_col_w] * 3)
    t_info_right.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), BrandColors.SLATE_50),
        ('BACKGROUND', (0, 1), (-1, 1), BrandColors.WHITE),
        ('BOX', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))

    # Outer two-panel card — explicit inner padding, no overflow
    t_info = Table(
        [[info_left, t_info_right]],
        colWidths=[PAGE_W * 0.58, PAGE_W * 0.42]
    )
    t_info.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, 0), BrandColors.SLATE_50),
        ('BACKGROUND', (1, 0), (1, 0), BrandColors.WHITE),
        ('BOX', (0, 0), (-1, -1), 0.75, BrandColors.SLATE_200),
        ('LINEAFTER', (0, 0), (0, -1), 0.75, BrandColors.SLATE_200),
        # Left cell padding
        ('LEFTPADDING', (0, 0), (0, 0), 14),
        ('RIGHTPADDING', (0, 0), (0, 0), 12),
        ('TOPPADDING', (0, 0), (0, 0), 12),
        ('BOTTOMPADDING', (0, 0), (0, 0), 12),
        # Right cell padding
        ('LEFTPADDING', (1, 0), (1, 0), 6),
        ('RIGHTPADDING', (1, 0), (1, 0), 6),
        ('TOPPADDING', (1, 0), (1, 0), 12),
        ('BOTTOMPADDING', (1, 0), (1, 0), 12),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 12))

    # ==============================================================
    # SECTION 3: ATTENDANCE KPI STRIP — Light
    # ==============================================================
    att_hdr_style = ParagraphStyle(
        'AttHdr', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=6.5, leading=9,
        textColor=BrandColors.SLATE_500, alignment=TA_CENTER,
    )
    att_val_style = ParagraphStyle(
        'AttVal', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=13, leading=15,
        textColor=BrandColors.SLATE_900, alignment=TA_CENTER,
    )
    att_sub_style = ParagraphStyle(
        'AttSub', parent=styles['Normal'],
        fontName='Helvetica', fontSize=6, leading=8,
        textColor=BrandColors.SLATE_400, alignment=TA_CENTER,
    )

    def att_kpi(value, unit, color_hex):
        return [
            Paragraph(f"<font color='{color_hex}'><b>{value}</b></font>", att_val_style),
            Paragraph(unit, att_sub_style),
        ]

    att_colors = {
        'attend': BrandColors.SUCCESS.hexval(),
        'absent': BrandColors.DANGER.hexval() if payroll.absent_days > 0 else BrandColors.SLATE_400.hexval(),
        'late': BrandColors.WARNING.hexval() if payroll.late_hours > 0 else BrandColors.SLATE_400.hexval(),
        'ot': BrandColors.BLUE.hexval() if payroll.overtime_hours > 0 else BrandColors.SLATE_400.hexval(),
        'leave': colors.HexColor('#7C3AED').hexval() if payroll.leave_days > 0 else BrandColors.SLATE_400.hexval(),
        'holiday': colors.HexColor('#D97706').hexval() if payroll.holiday_days > 0 else BrandColors.SLATE_400.hexval(),
    }

    att_labels = ['ATTENDANCE', 'ABSENCE', 'LATE HOURS', 'OVERTIME', 'LEAVES', 'HOLIDAYS']
    att_data_row = [
        att_kpi(payroll.attendance_days, 'days', att_colors['attend']),
        att_kpi(payroll.absent_days, 'days', att_colors['absent']),
        att_kpi(payroll.late_hours, 'hrs', att_colors['late']),
        att_kpi(payroll.overtime_hours, 'hrs', att_colors['ot']),
        att_kpi(payroll.leave_days, 'days', att_colors['leave']),
        att_kpi(payroll.holiday_days, 'days', att_colors['holiday']),
    ]

    kpi_col_w = PAGE_W / 6.0
    att_table_data = [
        [Paragraph(lbl, att_hdr_style) for lbl in att_labels],
        att_data_row,
    ]
    t_att = Table(att_table_data, colWidths=[kpi_col_w] * 6)
    t_att.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BrandColors.WHITE),
        ('BOX', (0, 0), (-1, -1), 0.75, BrandColors.SLATE_200),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
        ('TOPPADDING', (0, 1), (-1, 1), 4),
        ('BOTTOMPADDING', (0, 1), (-1, 1), 10),
        ('LINEABOVE', (0, 0), (-1, 0), 2, BrandColors.BLUE),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_att)
    story.append(Spacer(1, 12))

    # ==============================================================
    # SECTION 4: EARNINGS & DEDUCTIONS — Light, side-by-side
    # ==============================================================
    earn_hdr = ParagraphStyle(
        'EarnHdr', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7.5, leading=10,
        textColor=BrandColors.SUCCESS, alignment=TA_LEFT,
    )
    ded_hdr = ParagraphStyle(
        'DedHdr', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=7.5, leading=10,
        textColor=BrandColors.DANGER, alignment=TA_LEFT,
    )
    earn_label_s = ParagraphStyle('EarnLabel', parent=styles['Normal'],
                                   fontName='Helvetica-Bold', fontSize=8, leading=10,
                                   textColor=BrandColors.SLATE_900)
    earn_sub_s = ParagraphStyle('EarnSub', parent=styles['Normal'],
                                 fontName='Helvetica', fontSize=6.5, leading=8.5,
                                 textColor=BrandColors.SLATE_400)
    earn_amt_s = ParagraphStyle('EarnAmt', parent=styles['Normal'],
                                 fontName='Helvetica-Bold', fontSize=9, leading=11,
                                 textColor=BrandColors.SLATE_900, alignment=TA_RIGHT)
    ded_amt_s = ParagraphStyle('DedAmt', parent=styles['Normal'],
                                fontName='Helvetica-Bold', fontSize=9, leading=11,
                                textColor=BrandColors.SLATE_900, alignment=TA_RIGHT)
    total_label_s = ParagraphStyle('TotalLabel', parent=styles['Normal'],
                                    fontName='Helvetica-Bold', fontSize=8, leading=10,
                                    textColor=BrandColors.SLATE_900)
    total_earn_s = ParagraphStyle('TotalEarn', parent=styles['Normal'],
                                   fontName='Helvetica-Bold', fontSize=10, leading=12,
                                   textColor=BrandColors.SUCCESS, alignment=TA_RIGHT)
    total_ded_s = ParagraphStyle('TotalDed', parent=styles['Normal'],
                                  fontName='Helvetica-Bold', fontSize=10, leading=12,
                                  textColor=BrandColors.DANGER, alignment=TA_RIGHT)

    basic_label = "Basic Salary"
    if payroll.daily_salary:
        basic_label = f"Basic Salary  <font size=6 color='#94A3B8'>(${payroll.daily_salary:,.2f}/day)</font>"

    E_W = PAGE_W * 0.36
    EA_W = PAGE_W * 0.14
    SEP = 6
    D_W = E_W
    DA_W = EA_W

    amt_hdr_s = ParagraphStyle('AmtHdrS', parent=earn_hdr, alignment=TA_RIGHT, textColor=BrandColors.SUCCESS)
    amt_hdr_d = ParagraphStyle('AmtHdrD', parent=ded_hdr, alignment=TA_RIGHT, textColor=BrandColors.DANGER)

    breakdown_header = [
        Paragraph("  EARNINGS", earn_hdr),
        Paragraph("AMOUNT", amt_hdr_s),
        Paragraph("", earn_hdr),
        Paragraph("  DEDUCTIONS", ded_hdr),
        Paragraph("AMOUNT", amt_hdr_d),
    ]

    ot_sub = f"{payroll.overtime_hours} hrs recorded"
    ded_items = list(payroll.deduction_items.all()) if hasattr(payroll, 'deduction_items') else []
    if ded_items:
        items_summary = ", ".join(f"{it.category.replace('_', ' ').title()}: -${it.calculated_amount:,.2f}" for it in ded_items[:2])
        if len(ded_items) > 2:
            items_summary += f" (+{len(ded_items)-2} more)"
        other_ded_sub = items_summary
    else:
        other_ded_sub = "Unpaid absence, lateness"

    breakdown_rows = [
        [
            [Paragraph(basic_label, earn_label_s), Paragraph("Base monthly compensation", earn_sub_s)],
            Paragraph(f"${payroll.basic_salary:,.2f}", earn_amt_s),
            Paragraph("", styles['TableCell']),
            [Paragraph("Tax Withheld", earn_label_s), Paragraph("Salary Withholding Tax", earn_sub_s)],
            Paragraph(f"${payroll.tax:,.2f}", ded_amt_s),
        ],
        [
            [Paragraph("Allowances", earn_label_s), Paragraph("Transport, Meal &amp; Living", earn_sub_s)],
            Paragraph(f"${payroll.allowance:,.2f}", earn_amt_s),
            Paragraph("", styles['TableCell']),
            [Paragraph("NSSF Contribution", earn_label_s), Paragraph("Social Security Fund", earn_sub_s)],
            Paragraph(f"${payroll.nssf:,.2f}", ded_amt_s),
        ],
        [
            [Paragraph("Overtime Compensation", earn_label_s), Paragraph(ot_sub, earn_sub_s)],
            Paragraph(f"${payroll.overtime:,.2f}", earn_amt_s),
            Paragraph("", styles['TableCell']),
            [Paragraph("Other Deductions", earn_label_s), Paragraph(other_ded_sub, earn_sub_s)],
            Paragraph(f"${payroll.other_deduction:,.2f}", ded_amt_s),
        ],
        [
            [Paragraph("Bonus &amp; Incentives", earn_label_s), Paragraph("Performance &amp; Rewards", earn_sub_s)],
            Paragraph(f"${payroll.bonus:,.2f}", earn_amt_s),
            Paragraph("", styles['TableCell']),
            Paragraph("—", earn_label_s),
            Paragraph("—", ded_amt_s),
        ],
    ]

    total_row = [
        Paragraph("TOTAL GROSS EARNINGS", total_label_s),
        Paragraph(f"${payroll.gross_salary:,.2f}", total_earn_s),
        Paragraph("", styles['TableCell']),
        Paragraph("TOTAL DEDUCTIONS", total_label_s),
        Paragraph(f"-${payroll.total_deduction:,.2f}", total_ded_s),
    ]

    all_breakdown = [breakdown_header] + breakdown_rows + [total_row]
    col_widths = [E_W, EA_W, SEP, D_W, DA_W]

    t_break = Table(all_breakdown, colWidths=col_widths)
    t_break.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (1, 0), BrandColors.SUCCESS_PALE),
        ('BACKGROUND', (2, 0), (2, 0), BrandColors.WHITE),
        ('BACKGROUND', (3, 0), (4, 0), BrandColors.DANGER_PALE),
        ('BACKGROUND', (0, 1), (1, -2), BrandColors.WHITE),
        ('BACKGROUND', (3, 1), (4, -2), BrandColors.WHITE),
        ('BACKGROUND', (2, 1), (2, -1), BrandColors.WHITE),
        ('BACKGROUND', (0, -1), (1, -1), BrandColors.SUCCESS_PALE),
        ('BACKGROUND', (3, -1), (4, -1), BrandColors.DANGER_PALE),
        ('LINEABOVE', (0, -1), (1, -1), 1.2, BrandColors.SUCCESS),
        ('LINEABOVE', (3, -1), (4, -1), 1.2, BrandColors.DANGER),
        ('BOX', (0, 0), (1, -1), 0.5, BrandColors.SLATE_200),
        ('BOX', (3, 0), (4, -1), 0.5, BrandColors.SLATE_200),
        ('LINEBELOW', (0, 0), (1, 0), 0.6, BrandColors.SUCCESS),
        ('LINEBELOW', (3, 0), (4, 0), 0.6, BrandColors.DANGER),
        ('INNERGRID', (0, 1), (1, -2), 0.4, BrandColors.SLATE_100),
        ('INNERGRID', (3, 1), (4, -2), 0.4, BrandColors.SLATE_100),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (2, 0), (2, -1), 0),
        ('RIGHTPADDING', (2, 0), (2, -1), 0),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_break)
    story.append(Spacer(1, 12))

    # ==============================================================
    # SECTION 5: NET PAY BANNER — Light blue card
    # ==============================================================
    net_label_s = ParagraphStyle('NetLabel', parent=styles['Normal'],
                                  fontName='Helvetica-Bold', fontSize=7, leading=9,
                                  textColor=BrandColors.SLATE_500, alignment=TA_CENTER)
    net_amt_s = ParagraphStyle('NetAmt', parent=styles['Normal'],
                                fontName='Helvetica-Bold', fontSize=24, leading=28,
                                textColor=BrandColors.BLUE, alignment=TA_CENTER)
    net_formula_s = ParagraphStyle('NetFormula', parent=styles['Normal'],
                                    fontName='Helvetica', fontSize=7, leading=9,
                                    textColor=BrandColors.SLATE_500, alignment=TA_CENTER)
    net_period_lbl_s = ParagraphStyle('NetPeriodLbl', parent=styles['Normal'],
                                       fontName='Helvetica-Bold', fontSize=7, leading=9,
                                       textColor=BrandColors.SLATE_500, alignment=TA_CENTER)
    net_period_s = ParagraphStyle('NetPeriod', parent=styles['Normal'],
                                   fontName='Helvetica-Bold', fontSize=13, leading=16,
                                   textColor=BrandColors.SLATE_900, alignment=TA_CENTER)
    net_sub_s = ParagraphStyle('NetSub', parent=styles['Normal'],
                                fontName='Helvetica', fontSize=6.5, leading=8.5,
                                textColor=BrandColors.SLATE_400, alignment=TA_CENTER)

    formula_display = f"Gross ${payroll.gross_salary:,.2f}  —  Deductions ${payroll.total_deduction:,.2f}"

    net_left = [
        Paragraph("NET TAKE-HOME PAY", net_label_s),
        Spacer(1, 4),
        Paragraph(f"${payroll.net_salary:,.2f}", net_amt_s),
        Spacer(1, 3),
        Paragraph(formula_display, net_formula_s),
    ]

    currency_str = "Currency: USD ($)  •  Official Disbursement"
    period_str = payroll.payroll_period.strftime('%B %Y')
    net_right = [
        Paragraph("PAY PERIOD", net_period_lbl_s),
        Spacer(1, 4),
        Paragraph(period_str, net_period_s),
        Spacer(1, 6),
        Paragraph(f"<font color='{status_color.hexval()}'><b>{status_val}</b></font>", net_sub_s),
        Spacer(1, 4),
        Paragraph(currency_str, net_sub_s),
    ]

    t_net = Table([[net_left, net_right]], colWidths=[PAGE_W * 0.55, PAGE_W * 0.45])
    t_net.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BrandColors.BLUE_PALE),
        ('BOX', (0, 0), (-1, -1), 0.75, BrandColors.SLATE_200),
        ('LINEBEFORE', (0, 0), (0, -1), 4, BrandColors.BLUE),
        ('LINEAFTER', (0, 0), (0, -1), 0.5, BrandColors.SLATE_200),
        ('TOPPADDING', (0, 0), (-1, -1), 16),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 16),
        ('LEFTPADDING', (0, 0), (-1, -1), 14),
        ('RIGHTPADDING', (0, 0), (-1, -1), 14),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_net)
    story.append(Spacer(1, 12))

    # ==============================================================
    # SECTION 6: Legal Notice
    # ==============================================================
    story.append(HRFlowable(width="100%", thickness=0.5, color=BrandColors.SLATE_200))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "<i>This is an official computer-generated salary document issued by "
        f"{html.escape(company_name)}. All statutory deductions and tax withholdings "
        "comply with applicable labor regulations. For discrepancies, contact HR within 5 business days.</i>",
        styles['Notice']
    ))

    # ==============================================================
    # SECTION 7: Signature Block
    # ==============================================================
    story.append(build_signature_block(page_width=PAGE_W))

    doc.build(
        story,
        canvasmaker=lambda *a, **kw: ProfessionalCanvas(
            *a, **kw,
            document_title=f"Payslip - {emp_name}",
            document_ref=f"PS-{payroll.id:06d}"
        )
    )
    return buffer.getvalue()


# ==============================================================================
# 5. GENERIC ATTENDANCE & PAYROLL LIST EXPORTERS
# ==============================================================================

def generate_attendance_list_excel(attendances, company_name="Payroll MGM", generated_by="Admin"):
    """Excel export for a flat list of attendance records."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Records"
    ws.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws,
        company_name=company_name,
        report_title="Company Attendance Records",
        subtitle=f"Total Records: {len(attendances)}",
        generated_by=generated_by,
        max_col=9
    )

    headers = ["#", "Date", "Employee Code", "Employee Name", "Department",
               "Check In", "Check Out", "Working Hours", "Status"]
    ws.row_dimensions[5].height = 24

    for c_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=5, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    for idx, a in enumerate(attendances, start=1):
        r_idx = 5 + idx
        emp = a.employee
        cin = a.check_in.strftime("%H:%M") if a.check_in else "-"
        cout = a.check_out.strftime("%H:%M") if a.check_out else "-"
        hours_str = str(a.working_hours).split('.')[0] if a.working_hours else "-"

        ws.row_dimensions[r_idx].height = 20
        row_vals = [
            (idx, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (a.date.strftime("%Y-%m-%d"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (cin, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (cout, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (hours_str, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (a.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    ws.freeze_panes = 'A6'
    auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_payroll_list_excel(payrolls, company_name="Payroll MGM", generated_by="Admin"):
    """Excel export for a flat list of payroll records with totals."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll Records"
    ws.views.sheetView[0].showGridLines = True

    add_excel_header_block(
        ws,
        company_name=company_name,
        report_title="Company Payroll Records",
        subtitle=f"Total Records: {len(payrolls)}",
        generated_by=generated_by,
        max_col=17
    )

    headers = [
        "#", "Period", "Employee Code", "Employee Name", "Department",
        "Basic Salary", "Daily Salary", "Attendance", "Absent", "Late (Hrs)",
        "Overtime (Hrs)", "Leave", "Holiday", "Gross Salary", "Total Deductions",
        "Net Salary", "Status"
    ]
    ws.row_dimensions[5].height = 24

    for c_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=5, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    for idx, p in enumerate(payrolls, start=1):
        r_idx = 5 + idx
        emp = p.employee
        ws.row_dimensions[r_idx].height = 20
        row_vals = [
            (idx, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (p.payroll_period.strftime("%b %Y"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (float(p.basic_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.daily_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (f"{p.attendance_days}d", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{p.absent_days}d", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{p.late_hours}h", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{p.overtime_hours}h", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{p.leave_days}d", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{p.holiday_days}d", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (float(p.gross_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.total_deduction), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(p.net_salary), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (p.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    if payrolls:
        last_r = 5 + len(payrolls) + 1
        ws.row_dimensions[last_r].height = 22
        tots = [
            ("Grand Total", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            (f"=SUM(F6:F{last_r-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            ("", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
            (f"=SUM(N6:N{last_r-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            (f"=SUM(O6:O{last_r-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            (f"=SUM(P6:P{last_r-1})", ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_TOTAL, ExcelTheme.FMT_CURRENCY),
            ("", ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_TOTAL, None),
        ]
        for c_idx, (v, align, font, fmt) in enumerate(tots, start=1):
            cell = ws.cell(row=last_r, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_TOTAL
            cell.fill = ExcelTheme.TOTALS_FILL
            if fmt:
                cell.number_format = fmt

    ws.freeze_panes = 'A6'
    auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()