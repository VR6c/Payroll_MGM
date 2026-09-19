import io
import datetime
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from reports.exporters import (
    NumberedCanvas, ExcelTheme, auto_fit_columns, add_excel_header_block,
    get_pdf_styles, build_signature_block
)


def generate_overtime_report_excel(report_data, company_name="Payroll MGM", generated_by="Admin"):
    """
    Generates an executive-grade Excel workbook for Overtime Summary & Details.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Overtime Report"
    ws.views.sheetView[0].showGridLines = True

    date_str = f"{report_data['from_date']} to {report_data['to_date']}"
    add_excel_header_block(
        ws,
        company_name=company_name,
        report_title="Overtime Summary & Details Report",
        subtitle=f"Period: {date_str} | Total Hours: {report_data['total_hours']} hrs | Total Pay: ${report_data['total_amount']:.2f}",
        generated_by=generated_by,
        max_col=10
    )

    # --- KPI Summary Block (Rows 5 to 6) ---
    ws.cell(row=5, column=1, value="TOTAL HOURS").font = ExcelTheme.FONT_SUBTITLE
    ws.cell(row=6, column=1, value=f"{report_data['total_hours']:.2f} hrs").font = ExcelTheme.FONT_BOLD

    ws.cell(row=5, column=3, value="TOTAL OVERTIME PAY").font = ExcelTheme.FONT_SUBTITLE
    ws.cell(row=6, column=3, value=float(report_data['total_amount'])).number_format = ExcelTheme.FMT_CURRENCY
    ws.cell(row=6, column=3).font = ExcelTheme.FONT_BOLD

    ws.cell(row=5, column=5, value="APPROVED SESSIONS").font = ExcelTheme.FONT_SUBTITLE
    ws.cell(row=6, column=5, value=report_data['approved_count']).font = ExcelTheme.FONT_BOLD

    ws.cell(row=5, column=7, value="PENDING APPROVALS").font = ExcelTheme.FONT_SUBTITLE
    ws.cell(row=6, column=7, value=report_data['pending_count']).font = ExcelTheme.FONT_BOLD

    # --- Section 1: Department Breakdown ---
    start_r = 8
    ws.cell(row=start_r, column=1, value="DEPARTMENT OVERTIME SUMMARY").font = ExcelTheme.FONT_TITLE
    start_r += 1

    dept_headers = ["#", "Department", "Employees", "Total Hours", "Total Overtime Pay", "% of Total Pay"]
    ws.row_dimensions[start_r].height = 22
    for c_idx, h in enumerate(dept_headers, start=1):
        c = ws.cell(row=start_r, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    total_pay = float(report_data['total_amount']) or 1.0
    for idx, d in enumerate(report_data['dept_breakdown'], start=1):
        r_idx = start_r + idx
        ws.row_dimensions[r_idx].height = 20
        d_amt = float(d['total_amount'])
        pct = d_amt / total_pay if total_pay > 0 else 0

        row_vals = [
            (idx, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (d['dept_name'], ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (d['emp_count'], ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (float(d['total_hours']), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, "#,##0.00"),
            (d_amt, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (pct, ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_PERCENT),
        ]
        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    # --- Section 2: Detailed Records ---
    cur_r = start_r + len(report_data['dept_breakdown']) + 3
    ws.cell(row=cur_r, column=1, value="OVERTIME ACTIVITY LOG (DETAILED RECORDS)").font = ExcelTheme.FONT_TITLE
    cur_r += 1

    detail_headers = ["#", "Date", "Employee Code", "Employee Name", "Department", "Policy / Type", "Hours", "Multiplier", "Hourly Rate", "Overtime Pay", "Status"]
    ws.row_dimensions[cur_r].height = 24
    for c_idx, h in enumerate(detail_headers, start=1):
        c = ws.cell(row=cur_r, column=c_idx, value=h)
        c.font = ExcelTheme.FONT_HEADER
        c.fill = ExcelTheme.NAVY_HEADER_FILL
        c.alignment = ExcelTheme.ALIGN_HEADER
        c.border = ExcelTheme.BORDER_THIN

    for idx, ot in enumerate(report_data['records'], start=1):
        r_idx = cur_r + idx
        emp = ot.employee
        ws.row_dimensions[r_idx].height = 20

        row_vals = [
            (idx, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (ot.date.strftime("%Y-%m-%d"), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (emp.employee_code, ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, None),
            (f"{emp.first_name} {emp.last_name}".strip(), ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_BOLD, None),
            (emp.department.name if emp.department else "-", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (ot.overtime_type.name if ot.overtime_type else f"Standard ({ot.rate_multiplier}x)", ExcelTheme.ALIGN_LEFT, ExcelTheme.FONT_REGULAR, None),
            (float(ot.hours), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, "#,##0.00"),
            (float(ot.rate_multiplier), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_REGULAR, "0.0x"),
            (float(ot.hourly_rate), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_REGULAR, ExcelTheme.FMT_CURRENCY),
            (float(ot.overtime_amount), ExcelTheme.ALIGN_RIGHT, ExcelTheme.FONT_BOLD, ExcelTheme.FMT_CURRENCY),
            (ot.get_status_display(), ExcelTheme.ALIGN_CENTER, ExcelTheme.FONT_BOLD, None),
        ]

        zebra = ExcelTheme.ZEBRA_FILL if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, (v, align, font, fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font
            cell.alignment = align
            cell.border = ExcelTheme.BORDER_THIN
            cell.fill = zebra
            if fmt:
                cell.number_format = fmt

    auto_fit_columns(ws, min_width=12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_overtime_report_pdf(report_data, company_name="Payroll MGM", generated_by="Admin"):
    """
    Generates a PDF document for Overtime Summary & Details in Landscape format.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = get_pdf_styles()

    # Safely register all custom report styles
    if 'HeaderCompany' not in styles:
        styles.add(ParagraphStyle(
            name='HeaderCompany',
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#1E3A8A'),
        ))
    if 'DocTitle' not in styles:
        styles.add(ParagraphStyle(
            name='DocTitle',
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=3,
        ))
    if 'SubheaderCell' not in styles:
        styles.add(ParagraphStyle(
            name='SubheaderCell',
            fontName='Helvetica',
            fontSize=7.5,
            leading=10,
            textColor=colors.HexColor('#64748B'),
            alignment=1,
        ))
    if 'TitleStat' not in styles:
        styles.add(ParagraphStyle(
            name='TitleStat',
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#0F172A'),
            alignment=1,
        ))
    if 'TitleStatSuccess' not in styles:
        styles.add(ParagraphStyle(
            name='TitleStatSuccess',
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#059669'),
            alignment=1,
        ))
    if 'TitleStatWarning' not in styles:
        styles.add(ParagraphStyle(
            name='TitleStatWarning',
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#D97706'),
            alignment=1,
        ))
    if 'SectionTitle' not in styles:
        styles.add(ParagraphStyle(
            name='SectionTitle',
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=15,
            textColor=colors.HexColor('#1E293B'),
            spaceBefore=10,
            spaceAfter=6,
        ))
    if 'TableCellBoldRight' not in styles:
        styles.add(ParagraphStyle(
            name='TableCellBoldRight',
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor('#0F172A'),
            alignment=2,
        ))
    if 'TableHeaderCenter' not in styles:
        styles.add(ParagraphStyle(
            name='TableHeaderCenter',
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=11,
            textColor=colors.white,
            alignment=1,
        ))
    if 'ReportSubtitle' not in styles:
        styles.add(ParagraphStyle(
            name='ReportSubtitle',
            fontName='Helvetica',
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#64748B'),
            alignment=0,
        ))

    story = []

    # Title & Subtitle
    date_str = f"{report_data['from_date']} to {report_data['to_date']}"
    story.append(Paragraph(f"<b>{company_name.upper()}</b>", styles['HeaderCompany']))
    story.append(Paragraph("<b>OVERTIME SUMMARY & ANALYTICS REPORT</b>", styles['DocTitle']))
    story.append(Paragraph(f"Reporting Period: {date_str} | Generated By: {generated_by} on {datetime.date.today().strftime('%B %d, %Y')}", styles['ReportSubtitle']))
    story.append(Spacer(1, 12))

    # KPI Summary Cards in PDF Table
    kpi_data = [
        [
            Paragraph("<b>TOTAL HOURS</b>", styles['SubheaderCell']),
            Paragraph("<b>TOTAL OVERTIME PAY</b>", styles['SubheaderCell']),
            Paragraph("<b>APPROVED SESSIONS</b>", styles['SubheaderCell']),
            Paragraph("<b>PENDING REQUESTS</b>", styles['SubheaderCell'])
        ],
        [
            Paragraph(f"<b>{report_data['total_hours']:.2f} hrs</b>", styles['TitleStat']),
            Paragraph(f"<b>${report_data['total_amount']:.2f}</b>", styles['TitleStatSuccess']),
            Paragraph(f"<b>{report_data['approved_count']}</b>", styles['TitleStat']),
            Paragraph(f"<b>{report_data['pending_count']}</b>", styles['TitleStatWarning'])
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[185, 185, 185, 185])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F8FAFC')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#FFFFFF')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 16))

    # Section 1: Department Summary
    story.append(Paragraph("<b>DEPARTMENT OVERTIME BREAKDOWN</b>", styles['SectionTitle']))
    dept_table_data = [
        [
            Paragraph("<b>#</b>", styles['TableHeaderCenter']),
            Paragraph("<b>Department</b>", styles['TableHeader']),
            Paragraph("<b>Employees</b>", styles['TableHeaderCenter']),
            Paragraph("<b>Total Hours</b>", styles['TableHeaderRight']),
            Paragraph("<b>Total Pay ($)</b>", styles['TableHeaderRight']),
            Paragraph("<b>% Share</b>", styles['TableHeaderRight'])
        ]
    ]

    total_pay = float(report_data['total_amount']) or 1.0
    for idx, d in enumerate(report_data['dept_breakdown'], start=1):
        d_amt = float(d['total_amount'])
        pct = (d_amt / total_pay * 100) if total_pay > 0 else 0
        dept_table_data.append([
            Paragraph(str(idx), styles['TableCellCenter']),
            Paragraph(d['dept_name'], styles['TableCellBold']),
            Paragraph(str(d['emp_count']), styles['TableCellCenter']),
            Paragraph(f"{d['total_hours']:.2f} hrs", styles['TableCellRight']),
            Paragraph(f"${d_amt:.2f}", styles['TableCellBoldRight']),
            Paragraph(f"{pct:.1f}%", styles['TableCellRight']),
        ])

    dept_table = Table(dept_table_data, colWidths=[35, 230, 100, 120, 135, 120])
    dept_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(dept_table)
    story.append(Spacer(1, 16))

    # Section 2: Detailed Activity Log
    story.append(Paragraph("<b>OVERTIME DETAILED ACTIVITY LOG</b>", styles['SectionTitle']))
    rec_table_data = [
        [
            Paragraph("<b>#</b>", styles['TableHeaderCenter']),
            Paragraph("<b>Date</b>", styles['TableHeaderCenter']),
            Paragraph("<b>Employee</b>", styles['TableHeader']),
            Paragraph("<b>Department</b>", styles['TableHeader']),
            Paragraph("<b>Policy</b>", styles['TableHeader']),
            Paragraph("<b>Hours</b>", styles['TableHeaderRight']),
            Paragraph("<b>Rate</b>", styles['TableHeaderRight']),
            Paragraph("<b>Total Pay</b>", styles['TableHeaderRight']),
            Paragraph("<b>Status</b>", styles['TableHeaderCenter'])
        ]
    ]

    for idx, ot in enumerate(report_data['records'][:100], start=1):
        emp = ot.employee
        rec_table_data.append([
            Paragraph(str(idx), styles['TableCellCenter']),
            Paragraph(ot.date.strftime("%b %d, %Y"), styles['TableCellCenter']),
            Paragraph(f"{emp.employee_code} - {emp.first_name} {emp.last_name}", styles['TableCellBold']),
            Paragraph(emp.department.name if emp.department else "-", styles['TableCell']),
            Paragraph(ot.overtime_type.name if ot.overtime_type else "Standard", styles['TableCell']),
            Paragraph(f"{ot.hours:.2f} hrs", styles['TableCellRight']),
            Paragraph(f"${ot.hourly_rate:.2f}", styles['TableCellRight']),
            Paragraph(f"${ot.overtime_amount:.2f}", styles['TableCellBoldRight']),
            Paragraph(ot.get_status_display(), styles['TableCellCenter']),
        ])

    rec_table = Table(rec_table_data, colWidths=[30, 75, 175, 110, 130, 65, 60, 65, 60])
    rec_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(rec_table)

    # Signature Block
    story.append(build_signature_block(page_width=740))

    doc.build(story, canvasmaker=NumberedCanvas)
    return buf.getvalue()
